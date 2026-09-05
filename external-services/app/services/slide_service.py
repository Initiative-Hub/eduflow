import asyncio
import json
import logging
import os
import re
import shutil
import tempfile
import textwrap
import time
import zipfile
from functools import partial
from pathlib import Path
from typing import Any, Dict, List, Union
from fastapi.concurrency import run_in_threadpool
import slide_skills  # type: ignore
import slide_skills.svg_categories  # type: ignore
from app.deps import SLIDE_TEMPLATES_DIR
from app.schemas.slide_schema import RenderSlideReq

logger = logging.getLogger(__name__)

BASE_TEMPLATE_COLLECTION = "templates"

# A custom collection is pulled from S3 into SLIDE_TEMPLATES_DIR and swept once
# the request that pulled it is done. That directory is shared, so a request
# finishing first used to delete the layouts a slower one was still reading —
# surfacing as "[Errno 2] No such file or directory: .../<variant>.svg" partway
# through a deck. Requests register here and only the last one out sweeps.
_COLLECTION_USERS: Dict[str, int] = {}
_COLLECTION_USERS_LOCK = asyncio.Lock()


async def _acquire_collection(name: str) -> None:
    async with _COLLECTION_USERS_LOCK:
        _COLLECTION_USERS[name] = _COLLECTION_USERS.get(name, 0) + 1


async def _release_collection(name: str) -> bool:
    """Drop this request's claim; True when nobody else is using the cache."""
    async with _COLLECTION_USERS_LOCK:
        remaining = _COLLECTION_USERS.get(name, 1) - 1
        if remaining > 0:
            _COLLECTION_USERS[name] = remaining
            return False
        _COLLECTION_USERS.pop(name, None)
        return True


# The style inventory lists every key in both template buckets and then reads a
# collection.json per collection — a second of S3 work, repeated on every plan
# request and every time the picker opens, for an answer that only changes when
# somebody imports or deletes a template. Cached briefly so a burst of requests
# pays for it once; short enough that a new import shows up on its own.
_COLLECTIONS_TTL = float(os.environ.get("SLIDE_COLLECTIONS_CACHE_SECONDS", "60"))
# Files fetched at once when filling a collection's local cache from S3.
_DOWNLOAD_CONCURRENCY = max(1, int(os.environ.get("SLIDE_DOWNLOAD_CONCURRENCY", "16")))
_collections_cache: dict[str, Any] = {"at": 0.0, "value": None}
_COLLECTIONS_LOCK = asyncio.Lock()


def invalidate_collections_cache() -> None:
    """Drop the memoised style inventory after an import or a delete."""
    _collections_cache["at"] = 0.0
    _collections_cache["value"] = None


# "auto" detects brand templates whose designs live in the Slide Master layouts,
# "layouts" forces that reading, and "slides" extracts the deck's real slides.
TEMPLATE_IMPORT_SOURCES = {"auto", "layouts", "slides"}

# Rasterized category previews. Serving PNGs keeps the template picker light:
# inline SVG previews ship hundreds of KB of markup and build a live DOM tree
# per slide, which makes the dialog lag.
PREVIEW_FILE_NAME = "preview.png"
PREVIEW_WIDTH_PX = 1280

# Sample copy so previews look like real slides. `fill_svg` blanks any
# placeholder missing from this map instead of printing its raw name.
PREVIEW_SAMPLE_DATA: Dict[str, Any] = {
    "title": "Visual Learning Slide",
    "heading": "Concept Introduction",
    "subtitle": "A beautiful presentation design for courses and slides.",
    "presenter": "Presented by EduFlow",
    "author": "Presented by EduFlow",
    "kicker": "CHAPTER 1",
    "quote": '"Involve me and I learn."',
    "footer_note": "EduFlow Learning Platform",
    "body_text": (
        "Foundational concepts explained using modern slide layouts "
        "designed to keep students engaged."
    ),
    "left_col_title": "First Concept",
    "right_col_title": "Second Concept",
    "left_col_text": "Key details about the first concept side.",
    "right_col_text": "Comparison points on the second concept side.",
    "insight_text": "Engagement rose after switching to visual explanations.",
    # list_6 / points_6 variants exist, so these run to six: a short list left
    # the last rows of the six-row variants blank in their previews
    "bullets": [
        "Engaging detail or concept bullet point",
        "Supporting evidence for the concept",
        "A practical classroom example",
        "The mistake students make most often",
        "How to check the idea has landed",
        "Where it is used outside the classroom",
    ],
    "items": [
        "Introduction",
        "Core concepts",
        "Practice",
        "Summary",
        "Assessment",
        "Further reading",
    ],
    "steps": ["Prepare", "Explain", "Practise", "Review"],
    "summary_points": [
        "Key takeaway one",
        "Key takeaway two",
        "Key takeaway three",
        "Key takeaway four",
        "Key takeaway five",
        "Key takeaway six",
    ],
    "action_items": [
        "Read chapter 2",
        "Complete the worksheet",
        "Bring a question to the next session",
        "Review the summary notes",
        "Try the practice quiz",
    ],
}

# A slot named `heading.1` or `stat_2` is a different key from `heading` or
# `stat`, and fill_svg wipes any placeholder it has no value for. Cover slides,
# dividers, KPI cards, timelines and galleries were therefore rasterised empty,
# and the picker showed a blank rectangle for the very layouts that sell a
# style. Dotted and numbered keys pass through fill_svg untouched, so they are
# listed here explicitly alongside their scalar forms.
PREVIEW_SAMPLE_DATA.update(
    {
        "heading.1": "Concept",
        "heading.2": "Introduction",
        "heading.3": "for Learners",
        "sub_module_name": "MODULE ONE",
        "statement": "Great slides make the idea easier to hold on to.",
        "quote.1": "Involve me",
        "quote.2": "and I learn.",
        "author_or_source": "Benjamin Franklin",
        "media_image": "",
        **{f"body_text.{i}": text for i, text in enumerate(
            (
                "Foundational concepts explained with modern slide layouts.",
                "Each idea gets a claim and the reason it matters.",
                "Visual structure keeps a long explanation readable.",
                "Students leave with something they can act on.",
            ), 1)},
        **{f"left_col_text.{i}": t for i, t in enumerate(
            ("Key details about the first concept.",
             "What it looks like in practice.",
             "Where students usually get stuck.",
             "How to check understanding."), 1)},
        **{f"right_col_text.{i}": t for i, t in enumerate(
            ("Comparison points on the second concept.",
             "How the two differ in practice.",
             "When to prefer this approach.",
             "What it costs to adopt."), 1)},
        **{f"stat_{i}": v for i, v in enumerate(("87%", "3x", "1,240"), 1)},
        **{f"label_{i}": v for i, v in enumerate(
            ("of students stayed engaged",
             "faster to review before an exam",
             "lessons built this term"), 1)},
        **{f"caption_{i}": v for i, v in enumerate(
            ("A worked example from the lesson.",
             "The same idea shown as a diagram.",
             "Students applying it in class."), 1)},
        **{f"image_{i}": "" for i in range(1, 4)},
        **{f"date_{i}": v for i, v in enumerate(
            ("Week 1", "Week 3", "Week 6", "Week 9", "Week 11", "Week 12"), 1)},
        **{f"desc_{i}": v for i, v in enumerate(
            ("Introduce the core idea and why it matters.",
             "Work through an example together.",
             "Students practise with feedback.",
             "Review what stuck and what did not.",
             "Apply it to a new problem.",
             "Consolidate before assessment."), 1)},
        **{f"step_{i}": v for i, v in enumerate(
            ("Prepare the material and the question you want answered.",
             "Explain the idea with one concrete example.",
             "Let students practise while you watch for mistakes.",
             "Review the mistakes as a group.",
             "Set a short task that reuses the idea.",
             "Check understanding before moving on."), 1)},
        **{f"title_{i}": v for i, v in enumerate(
            ("Foundation", "Practice", "Feedback", "Mastery",
             "Transfer", "Assessment"), 1)},
        **{f"source_title_{i}": v for i, v in enumerate(
            ("Make It Stick: The Science of Successful Learning",
             "Rethinking Assessment in Higher Education",
             "Visible Learning for Teachers"), 1)},
        **{f"source_url_{i}": f"https://example.edu/library/source-{i}"
           for i in range(1, 4)},
    }
)

DEFAULT_COLLECTIONS = {
    "templates",
    "default",
    "starter",
    "neon_dark",
    "vintage",
    "pastel_pop",
    "illustrative_culture",
    "minimalist_gradient",
    "eduflow_light",
    "eduflow_purple",
    "rmit_official",
    "cultural_folk",
    "organic_streets",
    "green_environment_care",
    "startup_neon_pitch",
    "professional_focus",
}

STANDARD_LAYOUT_TYPES = [
    "TITLE_SLIDE",
    "AGENDA_OUTLINE",
    "SECTION_HEADER",
    "TITLE_BULLETS",
    "TWO_COLUMN_SPLIT",
    "BIG_QUOTE_TAKEAWAY",
    "KPI_BIG_NUMBER",
    "CHART_INSIGHT",
    "DATA_TABLE",
    "MEDIA_TEXT",
    "TIMELINE_MILESTONES",
    "STEP_BY_STEP",
    "CONCLUSION_SUMMARY",
    "CALL_TO_ACTION",
    "QA_CONTACT",
    "REFERENCES_LIST",
    "STATEMENT_IMAGE",
    "PYRAMID_LEVELS",
    "FUNNEL_STAGES",
    "PROCESS_ARROWS",
    "CIRCLE_CYCLE",
]

CATEGORY_ALIASES = {
    "CHART_SLIDE": "CHART_INSIGHT",
    "TABLE_SLIDE": "DATA_TABLE",
    "KPI_BIG_NUMBERS": "KPI_BIG_NUMBER",
    "IMAGE_TEXT": "MEDIA_TEXT",
    "AGENDA_AND_OUTLINE": "AGENDA_OUTLINE",
    "TITLE_AND_BULLETS": "TITLE_BULLETS",
}

# A customer deck names its layouts for its own designers — "Chart slide",
# "Content_option 2", "Section divider_option 1" — while the planner speaks the
# standard vocabulary. Nothing matched, so every category was backfilled from
# the base library and the deck came out in a foreign style even though the
# template had the layout all along. These keywords recognise an equivalent
# local layout so the customer's own design is used instead.
CATEGORY_EQUIVALENTS: Dict[str, tuple] = {
    "CHART_INSIGHT": ("chartslide", "chart", "graph", "barchart", "piechart"),
    "DATA_TABLE": ("tableslide", "table", "grid", "matrix"),
    "MEDIA_TEXT": ("contentwithimage", "imagetext", "picturetext", "imageandtext"),
    "IMAGE_GALLERY": ("imagegallery", "images", "gallery", "photogrid"),
    "TITLE_SLIDE": ("titleoption", "titleonly", "titleslide", "cover", "opening"),
    "SECTION_HEADER": ("sectiondivider", "section", "divider", "chapter"),
    "TITLE_BULLETS": ("titlebullets", "contentoption", "content", "bullet", "body"),
    "AGENDA_OUTLINE": ("agenda", "outline", "overview", "tableofcontents"),
    "CONCLUSION_SUMMARY": ("conclusion", "summary", "takeaway", "keypoints"),
    "QA_CONTACT": ("qacontact", "endslide", "thankyou", "question", "contact"),
    "KPI_BIG_NUMBER": ("kpi", "bignumber", "metric", "statistic"),
    "BIG_QUOTE_TAKEAWAY": ("quote", "testimonial"),
    "TIMELINE_MILESTONES": ("timeline", "milestone", "roadmap"),
    "STEP_BY_STEP": ("stepbystep", "process", "howto", "step"),
    "TWO_COLUMN_SPLIT": ("twocolumn", "comparison", "versus", "split"),
    "CALL_TO_ACTION": ("calltoaction", "nextstep", "cta"),
    "STATEMENT_IMAGE": ("statement", "hero", "impact"),
    "REFERENCES_LIST": ("reference", "source", "citation", "bibliography"),
}


def _canon_category(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def _best_standard_category(folder_name: str) -> tuple:
    """(category, keyword_length) this layout folder most specifically matches.

    Scored by the LONGEST matching keyword so "Content with image_option 1"
    lands on MEDIA_TEXT ("contentwithimage") rather than TITLE_BULLETS
    ("content"), which merely shares a prefix.
    """
    canon = _canon_category(folder_name)
    best, best_len = None, 0
    for category, keywords in CATEGORY_EQUIVALENTS.items():
        for keyword in keywords:
            if keyword in canon and len(keyword) > best_len:
                best, best_len = category, len(keyword)
    return best, best_len


def _find_equivalent_layout(library_dir: Path, category: str) -> Path | None:
    """A local layout folder that serves `category` under a different name."""
    matches = [
        child
        for child in sorted(library_dir.iterdir())
        if child.is_dir()
        and any(child.glob("*.svg"))
        and _best_standard_category(child.name)[0] == category
    ]
    if not matches:
        return None
    # Shortest name first: "Chart slide" over "Chart slide with commentary".
    return sorted(matches, key=lambda p: (len(p.name), p.name))[0]


CATEGORY_METADATA_FIELDS = (
    "description",
    "when_to_use",
    "prompt_hint",
    "content_guidance",
)


def _filter_category_metadata(raw: Any) -> Dict[str, Any]:
    if not isinstance(raw, dict):
        return {}
    return {
        key: value
        for key, value in raw.items()
        if key in CATEGORY_METADATA_FIELDS and value not in (None, "", [])
    }


def _read_category_metadata(meta_path: Path) -> Dict[str, Any]:
    if not meta_path.exists():
        return {}
    try:
        raw = json.loads(meta_path.read_text(encoding="utf-8"))
    except Exception as exc:
        logger.warning(f"Could not read category metadata from {meta_path}: {exc}")
        return {}

    return _filter_category_metadata(raw)


def _read_collection_category_metadata(library_dir: Path) -> Dict[str, Dict[str, Any]]:
    meta_path = library_dir / "collection.json"
    if not meta_path.exists():
        return {}

    try:
        raw = json.loads(meta_path.read_text(encoding="utf-8"))
    except Exception as exc:
        logger.warning(f"Could not read collection metadata from {meta_path}: {exc}")
        return {}

    categories = raw.get("categories")
    if not isinstance(categories, dict):
        return {}

    metadata: Dict[str, Dict[str, Any]] = {}
    for category, category_raw in categories.items():
        filtered = _filter_category_metadata(category_raw)
        if filtered:
            metadata[str(category)] = filtered

    return metadata


def _read_category_capacity(library_dir: Path | None) -> Dict[str, Dict[str, int]]:
    """Text/image slot counts per category, so planners avoid sending body copy
    to a layout that only has room for a title.

    Extracted brand templates are often sparse: a divider may expose a single
    `title` slot, and any extra binding is dropped when the deck is filled.
    """
    if not library_dir or not library_dir.exists():
        return {}

    try:
        library = slide_skills.scan_template_library(str(library_dir))
    except Exception as error:
        logger.warning(f"Could not read template capacity for {library_dir}: {error}")
        return {}

    capacity: Dict[str, Dict[str, int]] = {}
    for entry in library.category_map():
        variants = entry.get("variants") or []
        if not variants:
            continue
        # Report the roomiest variant: that is what the planner can rely on.
        capacity[entry["category"]] = {
            "text_slots": max(int(v.get("text_slots") or 0) for v in variants),
            "image_slots": max(int(v.get("image_slots") or 0) for v in variants),
            "capacity": max(int(v.get("capacity") or 0) for v in variants),
        }

    return capacity


def _build_category_metadata(
    library_dir: Path | None, categories: set[str]
) -> Dict[str, Dict[str, Any]]:
    metadata: Dict[str, Dict[str, Any]] = {}
    collection_metadata = (
        _read_collection_category_metadata(library_dir) if library_dir else {}
    )
    capacity = _read_category_capacity(library_dir)

    for category in sorted(categories):
        if not library_dir:
            continue
        merged = dict(collection_metadata.get(category, {}))
        category_metadata = _read_category_metadata(
            library_dir / category / "category.json"
        )
        for key, value in category_metadata.items():
            merged.setdefault(key, value)
        if category in capacity:
            merged.update(capacity[category])
        if merged:
            metadata[category] = merged

    return metadata


# Monkeypatch select_and_fill_slide to enforce outline bindings
_orig_select_and_fill_slide = slide_skills.svg_categories.select_and_fill_slide


def custom_select_and_fill_slide(variants, slide_content, **kwargs):
    res = _orig_select_and_fill_slide(variants, slide_content, **kwargs)
    if not res:
        return res

    variant = res.get("variant")
    texts = res.get("texts") or {}

    # Check both raw (unflattened) and flat bindings to support both multi-line
    # array placeholders (like "bullets", "items") and individual slot placeholders
    raw_bindings = slide_content.get("raw_bindings") or {}
    flat_bindings = slide_content.get("bindings") or {}

    for key in variant.placeholders:
        # Never clobber a value the library already resolved. It splits prose
        # across indexed placeholders (`title_2.1`, `.2`, `.3`) as a list, and
        # replacing that with the raw string leaves every indexed slot unfilled.
        # Since 0.2.42 bridges caller names onto slot names, planner values are
        # already present, so this only fills slots the library left empty.
        existing = texts.get(key)
        if existing not in (None, "", [], {}):
            continue
        if key in raw_bindings:
            texts[key] = raw_bindings[key]
        elif key in flat_bindings:
            texts[key] = flat_bindings[key]

    res["texts"] = texts
    return res


slide_skills.svg_categories.select_and_fill_slide = custom_select_and_fill_slide


# --- depth of the written copy -------------------------------------------------
# `select_and_fill_slide` writes the final text that lands on every slide, so it
# decides how substantial the deck reads — whatever the planner produced upstream
# is rewritten here to fit the template's slots.
#
# The library's own prompt gives that rewrite a ceiling ("NEVER exceed a
# placeholder's max_chars") and no floor. With nothing pushing the other way the
# model settles far under budget and pads the space with assertions: three cards
# each holding one slogan that restates its own heading. These rules add the
# missing floor and say what the words have to earn.
#
# Appended rather than replaced, so the library's hard rules — clipping, empty
# slots rendering as broken boxes, stat slots needing real figures, language
# matching — keep working and keep tracking upstream changes.
_DEPTH_RULES = """
DEPTH — what separates a usable slide from a hollow one:
- max_chars is a budget to SPEND, not merely a ceiling. For prose slots, land
  between 70% and 100% of it. Copy far under budget leaves the design looking
  empty and leaves the audience with nothing to take away.
- Explain, do not assert. Every prose slot must add something its own label and
  the slide title do not already say: the mechanism, a worked example, a figure,
  a consequence, or the condition under which it applies.
  Weak:   "Real exposure beats guesswork every time."
  Strong: "Two weeks shadowing a data team shows how much of the job is
           cleaning inputs, not modelling."
- Never restate the slide title, or the placeholder's own name, as its content.
- The supplied slide content is the source of truth. Carry its specifics through
  — names, numbers, dates, examples. Never trade a concrete detail for a generic
  phrase to save characters; drop a weaker clause instead.
- These rules govern prose slots only. Headings, labels, stats and figures stay
  as short as they are.
"""


def _with_depth_rules(system_prompt: str) -> str:
    """Insert the depth rules ahead of the prompt's output-format block."""
    if "DEPTH — what separates" in system_prompt:
        return system_prompt  # already applied; module re-imported
    marker = "Return ONLY JSON:"
    if marker in system_prompt:
        head, _, tail = system_prompt.partition(marker)
        return f"{head}{_DEPTH_RULES}\n{marker}{tail}"
    # Upstream reworded the output block — appending still reaches the model.
    return f"{system_prompt}\n{_DEPTH_RULES}"


slide_skills.svg_categories._SELECT_SYSTEM = _with_depth_rules(
    slide_skills.svg_categories._SELECT_SYSTEM
)

# --- style-aware image prompts -------------------------------------------------
# Generated images should match the collection's illustration style. The active
# collection's collection.json may define "image_style"; we set it here before
# generation and append it to every image prompt. If the planner supplied an
# explicit image_prompt_description binding, that takes precedence over the
# content-derived prompt.
_ACTIVE_IMAGE_STYLE: dict = {"hint": ""}
_orig_image_prompt = slide_skills.svg_categories._image_prompt


def _styled_image_prompt(slide, name, texts=None):
    bindings = (slide or {}).get("bindings") or {}
    desc = bindings.get("image_prompt_description")
    base = (
        f"A clear illustration of: {desc}. No text, no words, no letters."
        if isinstance(desc, str) and desc.strip()
        else _orig_image_prompt(slide, name, texts)
    )
    hint = _ACTIVE_IMAGE_STYLE.get("hint") or ""
    return f"{base} Art style: {hint}." if hint else base


slide_skills.svg_categories._image_prompt = _styled_image_prompt


def _set_image_style_for(library_dir) -> None:
    """Load the collection's image_style hint (if any) for prompt styling."""
    import json as _json

    _ACTIVE_IMAGE_STYLE["hint"] = ""
    try:
        meta = Path(library_dir) / "collection.json"
        if meta.exists():
            _ACTIVE_IMAGE_STYLE["hint"] = str(
                _json.loads(meta.read_text(encoding="utf-8")).get("image_style", "")
            )
    except Exception:
        pass


def _point_text(item: Any) -> str:
    """One list entry as a single line of slide copy.

    The planner may return a bare string, or a {title, description} pair when a
    point carries a claim AND its explanation. A slide slot is one line, so the
    pair is joined rather than dropped — losing the description is what made
    generated decks read as headline-only.
    """
    if isinstance(item, dict):
        title = str(item.get("title") or item.get("label") or "").strip()
        desc = str(item.get("description") or item.get("desc") or "").strip()
        if title and desc:
            return f"{title} — {desc}"
        return title or desc
    return str(item)


def flatten_slide_bindings(category: str, slide_title: str, bindings: dict) -> dict:
    flat = bindings.copy()

    # Always ensure slide title is mapped to heading for SVG template consistency
    if slide_title:
        flat["heading"] = slide_title
        # Multi-line heading for templates with a narrow heading panel
        # (e.g. TITLE_BULLETS/split_panel uses {{heading.1..3}}). Wrap so a
        # long title breaks across lines instead of overflowing. Widen until
        # the whole title fits in 3 lines (never drop words); fit_text_to_boxes
        # then shrinks any line still wider than the panel.
        wrap_width = max(16, -(-len(slide_title) // 3))  # ceil(len/3), min 16
        lines = textwrap.wrap(slide_title, width=wrap_width)
        while len(lines) > 3:
            wrap_width += 2
            lines = textwrap.wrap(slide_title, width=wrap_width)
        for idx, line in enumerate(lines, 1):
            flat[f"heading.{idx}"] = line

    # 1. Flatten AGENDA_OUTLINE: items -> items.1, items.2, etc.
    if "items" in bindings and isinstance(bindings["items"], list):
        for idx, item in enumerate(bindings["items"][:10], 1):
            flat[f"items.{idx}"] = _point_text(item)
            flat[f"item_{idx}"] = _point_text(item)

    # 2. Flatten TITLE_BULLETS: bullets -> bullets.1, bullets.2, etc.
    if "bullets" in bindings and isinstance(bindings["bullets"], list):
        for idx, item in enumerate(bindings["bullets"][:10], 1):
            flat[f"bullets.{idx}"] = _point_text(item)
            flat[f"bullet_{idx}"] = _point_text(item)

    # 3. Flatten TWO_COLUMN_SPLIT: left_col_text and right_col_text arrays
    if "left_col_text" in bindings and isinstance(bindings["left_col_text"], list):
        for idx, item in enumerate(bindings["left_col_text"][:10], 1):
            flat[f"left_col_text.{idx}"] = _point_text(item)
    if "right_col_text" in bindings and isinstance(bindings["right_col_text"], list):
        for idx, item in enumerate(bindings["right_col_text"][:10], 1):
            flat[f"right_col_text.{idx}"] = _point_text(item)

    # 4. Flatten BIG_QUOTE_TAKEAWAY: quote textwrap
    if "quote" in bindings and isinstance(bindings["quote"], str):
        # A fixed width guessed at a frame that does not exist. Chunks of 50
        # characters were far wider than the quote panel, so the renderer
        # wrapped each one a second time and a three-line quote sprawled into
        # nine, straight through the rule drawn beneath it — and `[:3]` threw
        # the tail away without saying so. Widen until the whole quote fits the
        # three lines the design has, keeping every word; fit_text_to_boxes then
        # shrinks the type for any line still wider than the panel.
        quote = bindings["quote"].strip()
        wrap_width = max(16, -(-len(quote) // 3))       # ceil(len / 3)
        lines = textwrap.wrap(quote, width=wrap_width)
        while len(lines) > 3:
            wrap_width += 2
            lines = textwrap.wrap(quote, width=wrap_width)
        for idx, line in enumerate(lines, 1):
            flat[f"quote.{idx}"] = line

    # 5. Flatten KPI_BIG_NUMBER: metrics -> stat_1, label_1, etc.
    if "metrics" in bindings and isinstance(bindings["metrics"], list):
        for idx, metric in enumerate(bindings["metrics"][:6], 1):
            if isinstance(metric, dict):
                flat[f"stat_{idx}"] = str(metric.get("value", ""))
                flat[f"label_{idx}"] = str(metric.get("label", ""))

    # 6. Flatten CHART_INSIGHT: chart_data & insight_text wrap
    if "chart_data" in bindings and isinstance(bindings["chart_data"], list):
        for idx, data_point in enumerate(bindings["chart_data"][:10], 1):
            if isinstance(data_point, dict):
                flat[f"chart_label_{idx}"] = str(data_point.get("label", ""))
                flat[f"chart_value_{idx}"] = str(
                    data_point.get("display_value", data_point.get("value", ""))
                )
    if "insight_text" in bindings and isinstance(bindings["insight_text"], str):
        wrapped = textwrap.wrap(bindings["insight_text"], width=45)
        for idx, line in enumerate(wrapped[:3], 1):
            flat[f"insight_text.{idx}"] = line

    # 7. Flatten DATA_TABLE: headers & rows
    if "headers" in bindings and isinstance(bindings["headers"], list):
        for idx, header in enumerate(bindings["headers"][:8], 1):
            flat[f"header_{idx}"] = str(header)
    if "rows" in bindings and isinstance(bindings["rows"], list):
        for r_idx, row in enumerate(bindings["rows"][:10], 1):
            if isinstance(row, list):
                for c_idx, val in enumerate(row[:8], 1):
                    flat[f"row_{r_idx}_{c_idx}"] = str(val)

    # 8. Flatten TIMELINE_MILESTONES: events date & desc wrap
    if "events" in bindings and isinstance(bindings["events"], list):
        for idx, ev in enumerate(bindings["events"][:10], 1):
            if isinstance(ev, dict):
                flat[f"date_{idx}"] = str(
                    ev.get("date_or_step", "") or ev.get("date", "")
                )
                desc = str(ev.get("description", ""))
                # Every timeline template declares a plain `desc_N`; only the
                # dotted sub-lines were emitted, so the description the planner
                # wrote reached no slot at all and the layout model invented a
                # replacement. The wrap is a hint for designs that split the
                # line themselves, and it no longer drops the tail.
                flat[f"desc_{idx}"] = desc
                for line_idx, line in enumerate(textwrap.wrap(desc, width=25), 1):
                    flat[f"desc_{idx}.{line_idx}"] = line

    # 8b. Flatten IMAGE_GALLERY: items -> label_N / caption_N. Nothing mapped
    # these before, so a gallery's own titles and captions were dropped and
    # rewritten downstream even though the plan already carried them.
    if category == "IMAGE_GALLERY" and isinstance(bindings.get("items"), list):
        for idx, item in enumerate(bindings["items"][:10], 1):
            if isinstance(item, dict):
                label = str(item.get("title", "") or item.get("label", ""))
                caption = str(item.get("description", "") or item.get("caption", ""))
            else:
                label, caption = str(item), ""
            if label:
                flat[f"label_{idx}"] = label
            if caption:
                flat[f"caption_{idx}"] = caption

    # 9. Flatten STEP_BY_STEP: steps -> step_1, step_2, etc.
    if "steps" in bindings and isinstance(bindings["steps"], list):
        for idx, step in enumerate(bindings["steps"][:10], 1):
            flat[f"step_{idx}"] = str(step)

    # 10. Flatten CONCLUSION_SUMMARY: summary_points -> summary_points.1, etc.
    if "summary_points" in bindings and isinstance(bindings["summary_points"], list):
        for idx, pt in enumerate(bindings["summary_points"][:10], 1):
            flat[f"summary_points.{idx}"] = str(pt)

    # 11. Flatten CALL_TO_ACTION: action_items -> action_items.1, etc.
    if "action_items" in bindings and isinstance(bindings["action_items"], list):
        for idx, item in enumerate(bindings["action_items"][:10], 1):
            flat[f"action_items.{idx}"] = _point_text(item)

    # 12. Flatten QA_CONTACT: footer_note
    if "footer_note" in bindings and isinstance(bindings["footer_note"], str):
        wrapped = textwrap.wrap(bindings["footer_note"], width=50)
        for idx, line in enumerate(wrapped[:2], 1):
            flat[f"footer_note.{idx}"] = line

    # 13. Flatten REFERENCES_LIST: sources -> source_title_1, source_url_1
    if "sources" in bindings and isinstance(bindings["sources"], list):
        for idx, src in enumerate(bindings["sources"][:10], 1):
            if isinstance(src, dict):
                flat[f"source_title_{idx}"] = str(src.get("title", ""))
                flat[f"source_url_{idx}"] = str(src.get("url", ""))

    # 13b. Flatten STATEMENT_IMAGE: statement -> statement.1..3 script lines
    if "statement" in bindings and isinstance(bindings["statement"], str):
        s_width = max(20, -(-len(bindings["statement"]) // 3))
        s_lines = textwrap.wrap(bindings["statement"], width=s_width)
        while len(s_lines) > 3:
            s_width += 2
            s_lines = textwrap.wrap(bindings["statement"], width=s_width)
        for idx, line in enumerate(s_lines, 1):
            flat[f"statement.{idx}"] = line

    # 14. Flatten diagram families (PYRAMID_LEVELS / FUNNEL_STAGES /
    #     PROCESS_ARROWS / CIRCLE_CYCLE): Array<{title, description}>
    #     -> title_1/desc_1 ... The per-count variants (tiers_3..5, stages_3..5,
    #     steps_3..5, phases_4..6) all use these same placeholder names, so one
    #     flattener covers every size.
    for diagram_key in ("levels", "stages", "process_steps", "phases"):
        if diagram_key in bindings and isinstance(bindings[diagram_key], list):
            for idx, item in enumerate(bindings[diagram_key][:10], 1):
                if isinstance(item, dict):
                    flat[f"title_{idx}"] = str(item.get("title", ""))
                    desc = str(item.get("description", ""))
                    flat[f"desc_{idx}"] = desc
                    # PROCESS_ARROWS descriptions are 2 short centered lines
                    wrapped = textwrap.wrap(desc, width=26)
                    for line_idx, line in enumerate(wrapped[:2], 1):
                        flat[f"desc_{idx}.{line_idx}"] = line
            break

    return flat


class SlideService:
    async def _ensure_collection_downloaded(self, collection: str | None) -> Path:
        if not collection:
            collection = "templates"

        col_path = Path(SLIDE_TEMPLATES_DIR) / collection
        has_local_layouts = col_path.exists() and any(col_path.glob("**/*.svg"))

        if collection.lower() in DEFAULT_COLLECTIONS:
            from app.deps import AWS_S3_DEFAULT_TEMPLATES_BUCKET as BUCKET_NAME
        else:
            from app.deps import AWS_S3_TEMPLATES_BUCKET as BUCKET_NAME

        from app.services.s3_service import (
            list_files_in_s3_prefix,
            download_file_from_s3,
        )

        s3_prefix = f"templates/{collection}/"
        s3_keys = await list_files_in_s3_prefix(s3_prefix, bucket_name=BUCKET_NAME)
        if not s3_keys:
            logger.warning(
                f"No files found in S3 bucket {BUCKET_NAME} for template collection '{collection}' at prefix '{s3_prefix}'"
            )
            if has_local_layouts:
                # Keep serving a locally cached collection when S3 is empty or
                # temporarily unreachable.
                return col_path
            # A missing collection must not degrade into an empty library dir,
            # which fails later with a confusing local-path error. Fall back to
            # the base layout collection instead.
            if collection != BASE_TEMPLATE_COLLECTION:
                logger.info(
                    f"Falling back to base template collection '{BASE_TEMPLATE_COLLECTION}'"
                )
                return await self._ensure_collection_downloaded(
                    BASE_TEMPLATE_COLLECTION
                )
            return col_path

        # Reconcile the local cache with S3 so previously truncated or partial
        # downloads are repaired instead of being treated as complete.
        pending: list[tuple[str, Path]] = []
        for key in s3_keys:
            relative = key[len(s3_prefix) :]
            if not relative or key.endswith("/"):
                continue
            dest = col_path / relative
            if not dest.exists():
                pending.append((key, dest))

        if not pending:
            return col_path

        logger.info(
            f"Downloading {len(pending)} file(s) for template collection "
            f"'{collection}' from S3 bucket {BUCKET_NAME}..."
        )
        col_path.mkdir(parents=True, exist_ok=True)
        for _, dest in pending:
            dest.parent.mkdir(parents=True, exist_ok=True)

        # One await per file meant one round-trip of latency per file, and a
        # collection is hundreds of small files. That is invisible against a
        # local MinIO and expensive against real S3 — and on a host with
        # ephemeral disk (Modal) every cold container pays it from scratch.
        semaphore = asyncio.Semaphore(_DOWNLOAD_CONCURRENCY)

        async def fetch(key: str, dest: Path) -> None:
            async with semaphore:
                await download_file_from_s3(key, dest, bucket_name=BUCKET_NAME)

        await asyncio.gather(*(fetch(key, dest) for key, dest in pending))
        return col_path

    async def get_categories(self) -> List[Dict[str, Any]]:
        # Always scan the downloaded default 'templates' library
        library_dir = await self._ensure_collection_downloaded("templates")
        library = await run_in_threadpool(
            slide_skills.scan_template_library, str(library_dir)
        )
        return library.category_map()

    def _template_bucket(self, collection: str) -> str:
        from app.deps import (
            AWS_S3_DEFAULT_TEMPLATES_BUCKET,
            AWS_S3_TEMPLATES_BUCKET,
        )

        return (
            AWS_S3_DEFAULT_TEMPLATES_BUCKET
            if collection.lower() in DEFAULT_COLLECTIONS
            else AWS_S3_TEMPLATES_BUCKET
        )

    @staticmethod
    def _render_category_preview(svg_path: Path, destination: Path) -> None:
        """Rasterizes one category layout to PNG, filled with sample copy."""
        import re

        import resvg_py
        from slide_skills.svg_collections import fit_and_reflow, infer_text_bounds

        svg = slide_skills.fill_svg(
            svg_path.read_text(encoding="utf-8"), PREVIEW_SAMPLE_DATA
        )
        # Filling alone is not what a real slide gets: deck generation also
        # wraps each block to its box and pushes later blocks clear. Without
        # those two steps the thumbnail drew every paragraph as one long
        # unwrapped line running straight through the card beside it, so the
        # picker misrepresented layouts that render fine in an actual deck.
        try:
            svg = fit_and_reflow(infer_text_bounds(svg))
        except Exception as error:  # a preview is not worth failing over
            logger.warning(f"Preview reflow failed for '{svg_path.name}': {error}")
        # Bare ampersands break the XML parser inside resvg.
        svg = re.sub(r"&(?!(?:[a-zA-Z0-9]+|#[0-9]+|#x[0-9a-fA-F]+);)", "&amp;", svg)
        destination.write_bytes(
            bytes(resvg_py.svg_to_bytes(svg_string=svg, width=PREVIEW_WIDTH_PX))
        )

    async def build_collection_previews(self, library_dir: Path) -> int:
        """Rasterizes a preview for every category in a local collection.

        Called at import time so opening the template picker never pays the
        rasterization cost, which takes several seconds for a large collection.
        """
        if not library_dir.exists():
            return 0

        rendered = 0
        for category_dir in sorted(
            child for child in library_dir.iterdir() if child.is_dir()
        ):
            variants = sorted(category_dir.glob("*.svg"))
            png_path = category_dir / PREVIEW_FILE_NAME
            if not variants or png_path.exists():
                continue
            try:
                await run_in_threadpool(
                    SlideService._render_category_preview, variants[0], png_path
                )
                rendered += 1
            except Exception as error:
                logger.warning(
                    f"Failed to rasterize preview for '{category_dir.name}': {error}"
                )

        return rendered

    async def get_template_previews(self, collection: str) -> Dict[str, Any]:
        """Returns one rasterized PNG preview per category, cached in S3.

        Previews are generated once and reused, so the template picker can load
        plain images instead of parsing and laying out full SVG markup.
        """
        from app.services.s3_service import (
            download_file_from_s3,
            list_files_in_s3_prefix,
            object_exists_in_s3,
            upload_file_to_s3,
        )

        bucket = self._template_bucket(collection)
        prefix = f"templates/{collection}/"

        # Fast path: previews already in S3 (generated at import time). Listing
        # keys avoids downloading the whole collection just to serve thumbnails.
        stored_keys = await list_files_in_s3_prefix(prefix, bucket_name=bucket)
        cached_keys = [
            key for key in stored_keys if key.endswith(f"/{PREVIEW_FILE_NAME}")
        ]
        categories_with_layouts = {
            key[len(prefix) :].rsplit("/", 1)[0]
            for key in stored_keys
            if key.endswith(".svg") and "/" in key[len(prefix) :]
        }
        categories_with_previews = {
            key[len(prefix) :].rsplit("/", 1)[0] for key in cached_keys
        }
        # Only skip the rebuild when every category has one, so a partially
        # generated collection is completed rather than shown with gaps.
        if cached_keys and categories_with_layouts <= categories_with_previews:
            return {
                "collection": collection,
                "bucket": bucket,
                "previews": [
                    {
                        "category": key[len(prefix) :].rsplit("/", 1)[0],
                        "variant": "standard",
                        "key": key,
                    }
                    for key in sorted(cached_keys)
                ],
            }

        library_dir = await self._ensure_collection_downloaded(collection)
        previews: List[Dict[str, str]] = []

        if not library_dir.exists():
            return {"collection": collection, "bucket": bucket, "previews": []}

        def render_png(svg_path: Path, destination: Path) -> None:
            SlideService._render_category_preview(svg_path, destination)

        for category_dir in sorted(
            child for child in library_dir.iterdir() if child.is_dir()
        ):
            variants = sorted(category_dir.glob("*.svg"))
            if not variants:
                continue

            variant = variants[0]
            object_key = (
                f"templates/{collection}/{category_dir.name}/{PREVIEW_FILE_NAME}"
            )
            png_path = category_dir / PREVIEW_FILE_NAME

            if not png_path.exists():
                cached = await object_exists_in_s3(
                    object_key, bucket_name=bucket
                ) and await download_file_from_s3(object_key, png_path, bucket)
                if not cached:
                    try:
                        await run_in_threadpool(render_png, variant, png_path)
                    except Exception as error:
                        logger.warning(
                            f"Failed to rasterize preview for "
                            f"'{collection}/{category_dir.name}': {error}"
                        )
                        continue

            # The key below is handed to the picker to sign and load, so the
            # object has to be there. Uploading only when the PNG had to be
            # rendered meant a collection whose previews already sat on local
            # disk — seeded into the image, or left over from an earlier run —
            # advertised keys that were never in the bucket, and every thumbnail
            # in the picker came back 404.
            if not await object_exists_in_s3(object_key, bucket_name=bucket):
                await upload_file_to_s3(png_path, object_key, bucket_name=bucket)

            previews.append(
                {
                    "category": category_dir.name,
                    "variant": variant.stem,
                    "key": object_key,
                }
            )

        return {"collection": collection, "bucket": bucket, "previews": previews}

    async def get_collection_categories(self, collection: str) -> Dict[str, Any]:
        library_dir = await self._ensure_collection_downloaded(collection)
        categories: set[str] = set()
        if library_dir.exists():
            for child in library_dir.iterdir():
                if child.is_dir() and child.name:
                    categories.add(child.name)

        is_custom = collection.lower() not in DEFAULT_COLLECTIONS

        if not categories:
            fallback_categories = set(STANDARD_LAYOUT_TYPES)
            fallback_metadata = _build_category_metadata(None, fallback_categories)
            return {
                "categories": STANDARD_LAYOUT_TYPES,
                "is_custom": is_custom,
                "metadata": fallback_metadata or None,
            }

        valid_categories = {cat for cat in categories if cat.isupper()}
        if not valid_categories:
            fallback_categories = set(STANDARD_LAYOUT_TYPES)
            fallback_metadata = _build_category_metadata(None, fallback_categories)
            return {
                "categories": STANDARD_LAYOUT_TYPES,
                "is_custom": is_custom,
                "metadata": fallback_metadata or None,
            }

        metadata = _build_category_metadata(library_dir, valid_categories)

        return {
            "categories": sorted(valid_categories),
            "is_custom": is_custom,
            "metadata": metadata or None,
        }

    async def get_collections(self) -> List[Dict[str, Any]]:
        cached = _collections_cache["value"]
        if cached is not None and time.monotonic() - _collections_cache["at"] < _COLLECTIONS_TTL:
            return cached
        async with _COLLECTIONS_LOCK:
            cached = _collections_cache["value"]
            if cached is not None and time.monotonic() - _collections_cache["at"] < _COLLECTIONS_TTL:
                return cached
            result = await self._load_collections()
            if result:
                _collections_cache["value"] = result
                _collections_cache["at"] = time.monotonic()
            return result

    async def _load_collections(self) -> List[Dict[str, Any]]:
        import json as _json
        import tempfile
        import asyncio
        from app.deps import AWS_S3_TEMPLATES_BUCKET, AWS_S3_DEFAULT_TEMPLATES_BUCKET
        from app.services.s3_service import list_files_in_s3_prefix

        categories = {
            "title_slide",
            "agenda_outline",
            "section_header",
            "title_bullets",
            "two_column_split",
            "big_quote_takeaway",
            "kpi_big_number",
            "chart_insight",
            "data_table",
            "media_text",
            "timeline_milestones",
            "step_by_step",
            "conclusion_summary",
            "call_to_action",
            "qa_contact",
            "references_list",
            "image_gallery",
            "kpi_big_numbers",
            "statement_image",
            "pyramid_levels",
            "funnel_stages",
            "process_arrows",
            "circle_cycle",
        }

        try:
            # 1. Fetch system template keys from AWS_S3_DEFAULT_TEMPLATES_BUCKET
            default_keys = await list_files_in_s3_prefix(
                "templates/", bucket_name=AWS_S3_DEFAULT_TEMPLATES_BUCKET
            )

            # 2. Fetch custom template keys from AWS_S3_TEMPLATES_BUCKET
            custom_keys = await list_files_in_s3_prefix(
                "templates/", bucket_name=AWS_S3_TEMPLATES_BUCKET
            )

            # Group keys by collection name along with their source bucket
            collections_files: Dict[str, Dict[str, Any]] = {}

            # Add default collections
            for key in default_keys:
                parts = key.split("/")
                if len(parts) >= 3 and parts[0] == "templates" and parts[1]:
                    col_name = parts[1]
                    if col_name not in collections_files:
                        collections_files[col_name] = {
                            "keys": [],
                            "bucket": AWS_S3_DEFAULT_TEMPLATES_BUCKET,
                        }
                    collections_files[col_name]["keys"].append(key)

            # Add custom collections
            for key in custom_keys:
                parts = key.split("/")
                if len(parts) >= 3 and parts[0] == "templates" and parts[1]:
                    col_name = parts[1]
                    if col_name not in collections_files:
                        collections_files[col_name] = {
                            "keys": [],
                            "bucket": AWS_S3_TEMPLATES_BUCKET,
                        }
                    collections_files[col_name]["keys"].append(key)

            async def get_description(name: str, keys: List[str], bucket: str) -> str:
                default_desc = f"Custom slide template collection '{name}'."
                json_key = f"templates/{name}/collection.json"
                if json_key not in keys:
                    well_known = {
                        "vintage": "A classic, retro style with warm tones and elegant typography.",
                        "pastel_pop": "A vibrant and playful theme featuring soft pastel colors.",
                        "starter": "Standard starter templates for clean presentation designs.",
                        "neon_dark": "A modern, high-contrast dark theme with glowing neon accents.",
                        "illustrative_culture": "Warm cream paper, hand-drawn buildings & clouds, Yogyakarta street aesthetic, sage green accents.",
                        "minimalist_gradient": "Sleek dark theme with electric royal blue and violet gradient glows, crisp geometric typography, and ambient grid lines.",
                        "organic_streets": "Organic illustration style: cream paper, plum script headlines, golden sun discs, slate and terracotta blobs, line-art European skylines.",
                        "eduflow_light": "Soft white canvas with violet accents and crisp bordered cards, from the EduFlow light theme - printed handouts, lectures projected in a bright room, and any deck that should look like EduFlow without going dark.",
                        "eduflow_purple": "Deep slate canvas with violet accents, soft bordered cards and an ambient glow — the EduFlow platform's own look. Course material, product walkthroughs, internal training, onboarding, and any deck that should feel native to the product it was made in.",
                        "rmit_official": "The official RMIT University brand template: navy grounds, RMIT red rules and headings, Arial throughout, and the university lock-up on the cover and closing slides — lectures, course material, research talks, student presentations, and anything that has to look like it came from RMIT.",
                        "cultural_folk": "Rich cultural folk style: warm plum night sky over a sand earth strip, arch and temple shapes, radiant sun badges, festival bunting and stitched lines in terracotta, gold, dusty blue and rose.",
                        "green_environment_care": "Modern environmental care style: cream paper, deep forest-green condensed headlines, lush nature photography, sage botanical ornaments, halftone texture, and conservation editorial layouts.",
                        "startup_neon_pitch": "Black startup pitch style with bold white typography, electric blue and violet light trails, glossy gradient pills, contact-footer details, and high-contrast business layouts.",
                        "professional_focus": "Calm executive presentation style with deep navy structure, precise teal signals, warm brass emphasis, generous whitespace, and business-ready editorial layouts.",
                    }
                    return well_known.get(name.lower(), default_desc)

                try:
                    with tempfile.TemporaryDirectory() as tmpdir:
                        temp_file = Path(tmpdir) / f"{name}_collection.json"
                        from app.services.s3_service import download_file_from_s3

                        downloaded = await download_file_from_s3(
                            json_key, temp_file, bucket_name=bucket
                        )
                        if downloaded:
                            content = temp_file.read_text(encoding="utf-8")
                            return _json.loads(content).get("description", default_desc)
                except Exception as e:
                    logger.warning(
                        f"Could not read collection.json for {name} from S3: {e}"
                    )
                return default_desc

            # Process names that are not individual categories and not the base "templates" or "default" collections
            valid_collections = [
                name
                for name in collections_files.keys()
                if name.lower() not in categories
                and name.lower() not in ("templates", "default")
            ]

            tasks = [
                get_description(
                    name,
                    collections_files[name]["keys"],
                    collections_files[name]["bucket"],
                )
                for name in sorted(valid_collections)
            ]
            descriptions = await asyncio.gather(*tasks)

            return [
                {
                    "name": name,
                    "description": desc,
                    "is_custom": name.lower() not in DEFAULT_COLLECTIONS,
                }
                for name, desc in zip(sorted(valid_collections), descriptions)
            ]
        except Exception as e:
            logger.warning(f"Could not list S3 template collections: {e}")
            return []

    async def generate_deck(
        self,
        topic: str,
        collection: str,
        output_path: Union[str, Path],
        palette: Union[str, tuple, None] = "auto",
        language: str | None = None,
        animation: str = "rise",
    ) -> Dict[str, Any]:
        col_name = collection or "templates"
        await self._ensure_collection_downloaded(col_name)
        return await run_in_threadpool(
            slide_skills.generate_web_deck,
            col_name,
            topic,
            output_path,
            palette=palette,
            language=language,
            animation=animation,
            base_dir=SLIDE_TEMPLATES_DIR,
        )

    async def _ensure_categories_exist(
        self, library_dir: Path, requested_categories: set[str]
    ) -> None:
        """Ensure all requested slide categories exist in library_dir.
        If a custom collection is missing a category (e.g. CHART_SLIDE or CHART_INSIGHT),
        backfill it from the base system 'templates' library so generation never fails
        with [Errno 2] No such file or directory."""
        if not library_dir.exists() or not library_dir.is_dir():
            return

        missing = [
            cat for cat in requested_categories
            if cat and not any((library_dir / cat).glob("*.svg"))
        ]
        if not missing:
            return

        base_dir = await self._ensure_collection_downloaded(BASE_TEMPLATE_COLLECTION)
        for cat in missing:
            cat_dir = library_dir / cat
            local_equivalent = _find_equivalent_layout(library_dir, cat)
            if local_equivalent is not None:
                cat_dir.mkdir(parents=True, exist_ok=True)
                for item in local_equivalent.iterdir():
                    if item.is_file():
                        shutil.copy2(item, cat_dir / item.name)
                logger.info(
                    f"Served '{cat}' from this collection's own "
                    f"'{local_equivalent.name}' layout instead of backfilling"
                )
                continue

            cat_dir.mkdir(parents=True, exist_ok=True)
            target_source = CATEGORY_ALIASES.get(cat, cat)
            src_dir = base_dir / target_source
            if not (src_dir.exists() and any(src_dir.glob("*.svg"))):
                src_dir = base_dir / cat

            if src_dir.exists() and any(src_dir.glob("*.svg")):
                for item in src_dir.iterdir():
                    if item.is_file():
                        shutil.copy2(item, cat_dir / item.name)
                logger.info(
                    f"Backfilled missing category '{cat}' in '{library_dir.name}' "
                    f"from base template '{src_dir.name}'"
                )
            else:
                fallback_dirs = [
                    d for d in base_dir.iterdir() if d.is_dir() and any(d.glob("*.svg"))
                ] or [
                    d for d in library_dir.iterdir() if d.is_dir() and any(d.glob("*.svg"))
                ]
                if fallback_dirs:
                    for item in fallback_dirs[0].iterdir():
                        if item.is_file():
                            shutil.copy2(item, cat_dir / item.name)
                    logger.warning(
                        f"Backfilled missing category '{cat}' in '{library_dir.name}' using fallback '{fallback_dirs[0].name}'"
                    )

    async def generate_deck_from_plan(
        self,
        plan: Union[list, dict],
        output_path: Union[str, Path],
        palette: Union[str, tuple, None] = "auto",
        language: str | None = None,
        animation: str = "rise",
        title: str | None = None,
        images: bool = True,
        image_source: str = "ai",
        collection: str | None = None,
        research: bool = True,
    ) -> Dict[str, Any]:
        slides_list = []
        if isinstance(plan, dict):
            slides_list = plan.get("slides", [])
        elif isinstance(plan, list):
            slides_list = plan

        requested_categories: set[str] = set()
        for slide in slides_list:
            bindings = slide.get("bindings") or {}
            raw_cat = slide.get("category") or slide.get("layoutType") or ""
            category = CATEGORY_ALIASES.get(raw_cat, raw_cat)
            slide_title = slide.get("slideTitle") or ""

            # Ensure category field is populated for slide_skills resolver
            slide["category"] = category
            if category:
                requested_categories.add(category)

            # Diagram and list families come in per-count variants. Expose the item
            # count as talking_points so slide_skills' capacity shortlist picks the
            # variant whose level/slot count matches the content exactly.
            list_keys = (
                "levels",
                "stages",
                "process_steps",
                "phases",
                "items",
                "bullets",
                "steps",
                "summary_points",
                "action_items",
                "metrics",
                "events",
                "chart_data",
                "sources",
            )
            for diagram_key in list_keys:
                items = bindings.get(diagram_key)
                if isinstance(items, list) and items:
                    slide["talking_points"] = [
                        str(
                            it.get("title", "")
                            or it.get("label", "")
                            or next(iter(it.values()), "")
                        )
                        if isinstance(it, dict)
                        else str(it)
                        for it in items
                    ]
                    break

            # Expose raw bindings for monkeypatched custom_select_and_fill_slide
            import json

            slide["raw_bindings"] = json.loads(json.dumps(bindings))

            # Flatten bindings to map to flat SVG placeholders
            bindings = flatten_slide_bindings(category, slide_title, bindings)

            if "body_text" in bindings and isinstance(bindings["body_text"], str):
                bindings["body_text"] = textwrap.wrap(bindings["body_text"], width=50)

            slide["bindings"] = bindings

        library_dir = SLIDE_TEMPLATES_DIR
        temp_dir_context = None

        col_name = collection or BASE_TEMPLATE_COLLECTION
        # Claim the cache before the download, so a request finishing now cannot
        # sweep the directory out from under this one.
        await _acquire_collection(col_name)
        collection_path = await self._ensure_collection_downloaded(col_name)
        if not any(collection_path.glob("**/*.svg")):
            await _release_collection(col_name)
            raise ValueError(
                f"Template collection '{col_name}' has no .svg layouts available. "
                "Verify it exists under the 'templates/<collection>/' prefix in the "
                "configured template S3 bucket."
            )
        if collection_path.exists() and collection_path.is_dir():
            # Detect format: new extract_template_smart collections have
            # category subdirectories (TITLE_SLIDE/, AGENDA_OUTLINE/, etc.)
            # while legacy imports store flat SVG files at the root.
            has_category_subdirs = any(
                child.is_dir() for child in collection_path.iterdir()
            )

            if has_category_subdirs:
                # New format — the collection dir IS a valid library_dir,
                # so pass it directly without any manual remapping.
                library_dir = collection_path
                logger.info(
                    f"Using new-format collection '{col_name}' directly as library_dir"
                )
            else:
                # Legacy flat SVG format — map files to category dirs manually
                import json

                temp_dir_context = tempfile.TemporaryDirectory()
                temp_lib_dir = Path(temp_dir_context.name)

                categories_to_map = [
                    "TITLE_SLIDE",
                    "AGENDA_OUTLINE",
                    "SECTION_HEADER",
                    "TITLE_BULLETS",
                    "TWO_COLUMN_SPLIT",
                    "BIG_QUOTE_TAKEAWAY",
                    "KPI_BIG_NUMBER",
                    "CHART_INSIGHT",
                    "DATA_TABLE",
                    "MEDIA_TEXT",
                    "TIMELINE_MILESTONES",
                    "STEP_BY_STEP",
                    "CONCLUSION_SUMMARY",
                    "CALL_TO_ACTION",
                    "QA_CONTACT",
                    "REFERENCES_LIST",
                ]

                patterns = {
                    "TITLE_SLIDE": ["title", "slide_00", "slide_title"],
                    "AGENDA_OUTLINE": ["agenda", "outline", "slide_01"],
                    "SECTION_HEADER": ["section", "header", "divider"],
                    "TITLE_BULLETS": [
                        "bullets",
                        "bullet",
                        "list",
                        "slide_02",
                        "slide_03",
                        "slide_04",
                    ],
                    "TWO_COLUMN_SPLIT": [
                        "two_column",
                        "two_col",
                        "split",
                        "columns",
                        "comparison",
                    ],
                    "BIG_QUOTE_TAKEAWAY": [
                        "quote",
                        "takeaway",
                        "saying",
                        "citation",
                    ],
                    "KPI_BIG_NUMBER": ["kpi", "statistic", "number", "metrics"],
                    "CHART_INSIGHT": ["chart", "insight", "graph", "visualization"],
                    "DATA_TABLE": ["table", "data", "grid"],
                    "MEDIA_TEXT": ["media", "image", "photo", "picture"],
                    "TIMELINE_MILESTONES": ["timeline", "milestones", "history"],
                    "STEP_BY_STEP": ["step", "process", "flow", "stages"],
                    "CONCLUSION_SUMMARY": ["conclusion", "summary", "closing"],
                    "CALL_TO_ACTION": ["cta", "action", "signup"],
                    "QA_CONTACT": ["qa", "contact", "questions", "social"],
                    "REFERENCES_LIST": ["reference", "sources", "citations"],
                }

                svg_files = list(collection_path.glob("*.svg"))

                for idx, cat in enumerate(categories_to_map):
                    cat_dir = temp_lib_dir / cat
                    cat_dir.mkdir(parents=True, exist_ok=True)

                    matched_file = None
                    candidates = patterns.get(cat, [])
                    for svg_file in svg_files:
                        stem_lower = svg_file.stem.lower()
                        for cand in candidates:
                            if cand in stem_lower:
                                matched_file = svg_file
                                break
                        if matched_file:
                            break

                    if not matched_file:
                        for svg_file in svg_files:
                            stem_lower = svg_file.stem.lower()
                            if (
                                cat.lower() in stem_lower
                                or cat.replace("_", "").lower() in stem_lower
                            ):
                                matched_file = svg_file
                                break

                    # Cycle through custom SVGs so style stays consistent
                    if not matched_file and svg_files:
                        matched_file = svg_files[idx % len(svg_files)]

                    if matched_file:
                        shutil.copy2(matched_file, cat_dir / "variant_a.svg")
                        schema_json = matched_file.with_suffix(".schema.json")
                        if schema_json.exists():
                            shutil.copy2(schema_json, cat_dir / "variant_a.schema.json")

                    (cat_dir / "category.json").write_text(
                        json.dumps(
                            {
                                "description": f"{cat} category layout",
                                "variants": {"variant_a": "Default design"},
                            },
                            ensure_ascii=False,
                        )
                    )

                library_dir = temp_lib_dir

        if library_dir and requested_categories:
            await self._ensure_categories_exist(library_dir, requested_categories)

        try:
            actual_palette = None if palette == "auto" else palette
            _set_image_style_for(library_dir)  # style-matched AI image prompts
            res = await run_in_threadpool(
                slide_skills.generate_deck_from_plan,
                plan,
                library_dir,
                output_path,
                palette=actual_palette,
                language=language,
                animation=animation,
                images=images,
                image_source=image_source,
                title=title,
                research=research,
            )
        finally:
            if temp_dir_context:
                try:
                    temp_dir_context.cleanup()
                except Exception:
                    pass
            # Remove the S3-downloaded collection from local disk after use,
            # but only once every concurrent request has finished with it.
            last_user = await _release_collection(col_name) if col_name else True
            if col_name and col_name not in DEFAULT_COLLECTIONS and last_user:
                col_path = Path(SLIDE_TEMPLATES_DIR) / col_name
                if col_path.exists() and col_path.is_dir():
                    shutil.rmtree(col_path, ignore_errors=True)
                    logger.info(
                        f"Cleaned up downloaded S3 collection '{col_name}' after generation"
                    )

        # slide-skills reports plan data that had no matching slot. Log it so a
        # sparse template silently swallowing content is diagnosable.
        for warning in res.get("warnings") or []:
            if "dropped bindings" in warning:
                logger.warning(f"[{col_name}] {warning}")

        return res

    async def render_slide(self, req: RenderSlideReq) -> Dict[str, Any]:

        bindings: Dict[str, Any] = req.bindings.copy() if req.bindings else {}
        slide: Dict[str, Any] = {
            "category": req.layoutType,
            "slideTitle": req.slideTitle,
            "bindings": bindings,
        }
        category = req.layoutType
        slide_title = req.slideTitle

        slide["category"] = category

        # Diagram and list families come in per-count variants. Expose the item
        # count as talking_points so slide_skills' capacity shortlist picks the
        # variant whose level/slot count matches the content exactly.
        list_keys = (
            "levels",
            "stages",
            "process_steps",
            "phases",
            "items",
            "bullets",
            "steps",
            "summary_points",
            "action_items",
            "metrics",
            "events",
            "chart_data",
            "sources",
        )
        for diagram_key in list_keys:
            items = bindings.get(diagram_key)
            if isinstance(items, list) and items:
                slide["talking_points"] = [
                    str(
                        it.get("title", "")
                        or it.get("label", "")
                        or next(iter(it.values()), "")
                    )
                    if isinstance(it, dict)
                    else str(it)
                    for it in items
                ]
                break

        # Expose raw bindings for monkeypatched custom_select_and_fill_slide
        import json

        slide["raw_bindings"] = json.loads(json.dumps(bindings))

        bindings = flatten_slide_bindings(category, slide_title, bindings)

        if "body_text" in bindings and isinstance(bindings["body_text"], str):
            bindings["body_text"] = textwrap.wrap(bindings["body_text"], width=50)

        slide["bindings"] = bindings

        col_name = req.collection or "templates"
        await _acquire_collection(col_name)
        collection_path = await self._ensure_collection_downloaded(col_name)

        library_dir = SLIDE_TEMPLATES_DIR
        temp_dir_context = None

        if collection_path.exists() and collection_path.is_dir():
            has_category_subdirs = any(
                child.is_dir() for child in collection_path.iterdir()
            )
            if has_category_subdirs:
                library_dir = collection_path
            else:
                # Legacy flat format
                import json

                temp_dir_context = tempfile.TemporaryDirectory()
                temp_lib_dir = Path(temp_dir_context.name)

                categories_to_map = [
                    "TITLE_SLIDE",
                    "AGENDA_OUTLINE",
                    "SECTION_HEADER",
                    "TITLE_BULLETS",
                    "TWO_COLUMN_SPLIT",
                    "BIG_QUOTE_TAKEAWAY",
                    "KPI_BIG_NUMBER",
                    "CHART_INSIGHT",
                    "DATA_TABLE",
                    "MEDIA_TEXT",
                    "TIMELINE_MILESTONES",
                    "STEP_BY_STEP",
                    "CONCLUSION_SUMMARY",
                    "CALL_TO_ACTION",
                    "QA_CONTACT",
                    "REFERENCES_LIST",
                ]

                patterns = {
                    "TITLE_SLIDE": ["title", "slide_00", "slide_title"],
                    "AGENDA_OUTLINE": ["agenda", "outline", "slide_01"],
                    "SECTION_HEADER": ["section", "header", "slide_02"],
                    "TITLE_BULLETS": ["bullets", "bullet", "points", "slide_03"],
                    "TWO_COLUMN_SPLIT": ["split", "columns", "slide_04"],
                    "BIG_QUOTE_TAKEAWAY": ["quote", "takeaway", "slide_05"],
                    "KPI_BIG_NUMBER": ["kpi", "number", "metric", "slide_06"],
                    "CHART_INSIGHT": ["chart", "insight", "graph", "slide_07"],
                    "DATA_TABLE": ["table", "data_table", "slide_08"],
                    "MEDIA_TEXT": ["media", "image_text", "slide_09"],
                    "TIMELINE_MILESTONES": ["timeline", "milestone", "slide_10"],
                    "STEP_BY_STEP": ["step", "process_steps", "slide_11"],
                    "CONCLUSION_SUMMARY": ["conclusion", "summary", "slide_12"],
                    "CALL_TO_ACTION": ["cta", "action", "slide_13"],
                    "QA_CONTACT": ["qa", "contact", "slide_14"],
                    "REFERENCES_LIST": ["references", "source", "slide_15"],
                }

                svg_files = list(collection_path.glob("*.svg"))
                for idx, cat in enumerate(categories_to_map):
                    cat_dir = temp_lib_dir / cat
                    cat_dir.mkdir(parents=True, exist_ok=True)

                    matched_file = None
                    for pattern in patterns.get(cat, []):
                        for svg_file in svg_files:
                            if pattern in svg_file.name.lower():
                                matched_file = svg_file
                                break
                        if matched_file:
                            break

                    if not matched_file:
                        for svg_file in svg_files:
                            stem_lower = svg_file.stem.lower()
                            if (
                                cat.lower() in stem_lower
                                or cat.replace("_", "").lower() in stem_lower
                            ):
                                matched_file = svg_file
                                break

                    if not matched_file and svg_files:
                        matched_file = svg_files[idx % len(svg_files)]

                    if matched_file:
                        shutil.copy2(matched_file, cat_dir / "variant_a.svg")
                        schema_json = matched_file.with_suffix(".schema.json")
                        if schema_json.exists():
                            shutil.copy2(schema_json, cat_dir / "variant_a.schema.json")

                    (cat_dir / "category.json").write_text(
                        json.dumps(
                            {
                                "description": f"{cat} category layout",
                                "variants": {"variant_a": "Default design"},
                            },
                            ensure_ascii=False,
                        )
                    )

                library_dir = temp_lib_dir

        if library_dir and category:
            await self._ensure_categories_exist(library_dir, {category})

        try:
            lib = await run_in_threadpool(
                slide_skills.scan_template_library, library_dir
            )
            key = lib.resolve(category)
            if key is None:
                raise ValueError(
                    f"No category matches {category!r} in template library"
                )

            mapping = {}
            target = (
                slide_skills.PRESETS.get(req.palette)
                if isinstance(req.palette, str)
                else req.palette
            )
            if target is not None:
                from slide_skills.svg_categories import _library_palette

                mapping = slide_skills.auto_map_palette(_library_palette(lib), target)

            result = await run_in_threadpool(
                slide_skills.svg_categories.select_and_fill_slide,
                lib.categories[key],
                slide,
            )
            logger.info(
                f"===> RENDER_SLIDE RESULT: {result.get('variant').name if result else None}"
            )
            logger.info(
                f"===> RENDER_SLIDE INPUT TALKING POINTS: {slide.get('talking_points')}"
            )
            logger.info(f"===> RENDER_SLIDE INPUT BINDINGS: {slide.get('bindings')}")
            if not result:
                raise ValueError(f"No variant selected for layout {category}")

            variant = result["variant"]
            svg = Path(variant.path).read_text(encoding="utf-8")

            # Data-slot renderers need a theme-aware ink colour for labels and
            # axes. Reuse the heading colour so dark collections receive light
            # chart text while light collections keep their normal dark ink.
            slot_ink = "#1A1A1A"
            heading_tag = re.search(
                r'<text\b[^>]*>\s*\{\{heading(?:\|[^}]*)?\}\}\s*</text>',
                svg,
                flags=re.IGNORECASE,
            )
            if heading_tag:
                heading_fill = re.search(
                    r'\bfill="(#[0-9A-Fa-f]{6})"', heading_tag.group(0)
                )
                if heading_fill:
                    slot_ink = heading_fill.group(1)
            slot_tag = re.search(
                r'<rect\b[^>]*\bdata-slot="(?:chart|table)"[^>]*/?>',
                svg,
                flags=re.IGNORECASE,
            )
            if slot_tag:
                explicit_ink = re.search(
                    r'\bdata-ink="(#[0-9A-Fa-f]{6})"', slot_tag.group(0)
                )
                if explicit_ink:
                    slot_ink = explicit_ink.group(1)

            from slide_skills.svg_categories import (
                prune_empty_groups,
                fill_svg,
                fit_text_to_boxes,
                reflow_text_blocks,
            )
            from slide_skills.svg_charts import fill_data_slots

            svg = prune_empty_groups(svg, result["texts"])

            texts = result["texts"]
            if 'data-w="' in svg:
                from slide_skills.svg_categories import _backfill_slots

                texts = _backfill_slots(texts, variant, slide)

            svg = fill_svg(svg, texts)
            raw_binds = slide.get("raw_bindings") or slide.get("bindings") or {}
            svg = fill_data_slots(svg, raw_binds, target, ink=slot_ink)
            svg = fit_text_to_boxes(svg)
            # push later paragraphs down so wrapped lines can't overlap
            svg = reflow_text_blocks(svg)

            if mapping:
                svg = slide_skills.retheme_svg(svg, mapping)

            return {"svg": svg}
        finally:
            if temp_dir_context:
                try:
                    temp_dir_context.cleanup()
                except Exception:
                    pass
            last_user = await _release_collection(col_name) if col_name else True
            if col_name and col_name not in DEFAULT_COLLECTIONS and last_user:
                col_path = Path(SLIDE_TEMPLATES_DIR) / col_name
                if col_path.exists() and col_path.is_dir():
                    shutil.rmtree(col_path, ignore_errors=True)
                    logger.info(
                        f"Cleaned up downloaded S3 collection '{col_name}' after rendering"
                    )

    async def import_template_collection(
        self,
        file_bytes: bytes,
        filename: str,
        name: str | None = None,
        source: str = "auto",
    ) -> Dict[str, Any]:
        if source not in TEMPLATE_IMPORT_SOURCES:
            raise ValueError(f"source must be one of {sorted(TEMPLATE_IMPORT_SOURCES)}")

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            collection_name = name or Path(filename).stem

            if filename.lower().endswith(".pptx"):
                # extract_template_smart maps slides to categories via AI and
                # writes a CATEGORY/standard.svg structure into library_dir.
                # `source` lets brand templates be read from the Slide Master's
                # layouts, where corporate designs actually live.
                pptx_path = temp_path / filename
                pptx_path.write_bytes(file_bytes)

                res = await run_in_threadpool(
                    partial(
                        slide_skills.extract_template_smart,
                        str(pptx_path),
                        collection_name,
                        library_dir=SLIDE_TEMPLATES_DIR,
                        use_ai=True,
                        source=source,
                    )
                )

            elif filename.lower().endswith(".zip"):
                zip_path = temp_path / "upload.zip"
                zip_path.write_bytes(file_bytes)

                extract_path = temp_path / "extracted"
                extract_path.mkdir()

                with zipfile.ZipFile(zip_path, "r") as zip_ref:
                    zip_ref.extractall(extract_path)

                # If there is a nested folder, use it
                subdirs = [x for x in extract_path.iterdir() if x.is_dir()]
                src_path = subdirs[0] if len(subdirs) == 1 else extract_path

                res = await run_in_threadpool(
                    slide_skills.import_collection,
                    src_path,
                    collection_name,
                    base_dir=SLIDE_TEMPLATES_DIR,
                    overwrite=True,
                )

            else:
                svg_path = temp_path / filename
                svg_path.write_bytes(file_bytes)

                res = await run_in_threadpool(
                    slide_skills.import_collection,
                    svg_path,
                    collection_name,
                    base_dir=SLIDE_TEMPLATES_DIR,
                    overwrite=True,
                )

            # Upload the imported templates to S3 bucket then remove local copy.
            # extract_template_smart writes to <collection_name>_template/ while
            # import_collection writes to <collection_name>/. Detect which was used.
            from app.deps import AWS_S3_TEMPLATES_BUCKET
            from app.services.s3_service import upload_file_to_s3

            dest_dir = Path(SLIDE_TEMPLATES_DIR) / collection_name
            # Check for the _template suffix variant produced by extract_template_smart
            dest_dir_with_suffix = (
                Path(SLIDE_TEMPLATES_DIR) / f"{collection_name}_template"
            )
            if not dest_dir.exists() and dest_dir_with_suffix.exists():
                dest_dir = dest_dir_with_suffix

            # Two layouts classified into the same category overwrite each other,
            # so surface the loss instead of silently importing fewer designs.
            warnings = list(res.get("warnings", [])) if isinstance(res, dict) else []
            if isinstance(res, dict) and res.get("categories"):
                classified = [
                    entry.get("category")
                    for entry in res["categories"]
                    if isinstance(entry, dict) and entry.get("category")
                ]
                folder_count = (
                    len([child for child in dest_dir.iterdir() if child.is_dir()])
                    if dest_dir.exists() and dest_dir.is_dir()
                    else 0
                )
                if folder_count and folder_count < len(classified):
                    message = (
                        f"{len(classified) - folder_count} of {len(classified)} layouts "
                        "shared a category name and were overwritten during import."
                    )
                    logger.warning(f"[{collection_name}] {message}")
                    warnings.append(message)
            if isinstance(res, dict) and warnings:
                res["warnings"] = warnings

            if dest_dir.exists() and dest_dir.is_dir():
                # Auto-create collection.json if it doesn't exist
                meta_file = dest_dir / "collection.json"
                if not meta_file.exists():
                    import json

                    try:
                        meta_file.write_text(
                            json.dumps(
                                {
                                    "name": collection_name,
                                    "description": f"Custom slide template collection '{collection_name}'.",
                                },
                                indent=2,
                                ensure_ascii=False,
                            ),
                            encoding="utf-8",
                        )
                        logger.info(
                            f"Created default collection.json for custom collection '{collection_name}'"
                        )
                    except Exception as e:
                        logger.warning(
                            f"Could not auto-create collection.json for '{collection_name}': {e}"
                        )

                async def upload_dir_to_s3(directory: Path, prefix: str) -> int:
                    uploaded_count = 0
                    for child in directory.iterdir():
                        if child.is_file():
                            uploaded = await upload_file_to_s3(
                                child,
                                f"{prefix}/{child.name}",
                                bucket_name=AWS_S3_TEMPLATES_BUCKET,
                            )
                            if not uploaded:
                                raise RuntimeError(
                                    f"Failed to upload template file '{child}' to S3 bucket '{AWS_S3_TEMPLATES_BUCKET}'"
                                )
                            uploaded_count += 1
                        elif child.is_dir():
                            uploaded_count += await upload_dir_to_s3(
                                child, f"{prefix}/{child.name}"
                            )
                    return uploaded_count

                # Rasterize category previews now so they upload with the
                # collection. Generating them on first open makes the template
                # picker spin for seconds on a large collection.
                rendered = await self.build_collection_previews(dest_dir)
                logger.info(
                    f"Rendered {rendered} preview image(s) for collection '{collection_name}'"
                )

                # Always use collection_name (without _template) as the S3 prefix
                # so get_collections and generate_deck_from_plan can find it consistently
                uploaded_count = await upload_dir_to_s3(
                    dest_dir, f"templates/{collection_name}"
                )
                if uploaded_count == 0:
                    raise RuntimeError(
                        f"No template files were uploaded to S3 for collection '{collection_name}'"
                    )

                # Remove local files after successful S3 upload
                shutil.rmtree(dest_dir, ignore_errors=True)
                logger.info(
                    f"Removed local template collection '{dest_dir.name}' after uploading {uploaded_count} files to S3"
                )

            # a new collection must appear in the picker now, not in a minute
            invalidate_collections_cache()
            return res

    async def generate_pptx(self, deck_id: str) -> Path:
        import base64
        import copy
        import io
        import re
        import traceback
        import xml.etree.ElementTree as ET

        import resvg_py
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu, Inches, Pt

        from app.deps import STORAGE_DIR
        from app.services.s3_service import (
            download_bytes_from_s3,
            download_file_from_s3,
        )

        # Always fetch latest HTML from S3 to make sure we have visual edits
        s3_key = f"slides/{deck_id}.html"
        local_html_path = STORAGE_DIR / f"{deck_id}_latest.html"

        # Try downloading from S3
        downloaded = await download_file_from_s3(s3_key, local_html_path)
        if not downloaded:
            # Fallback: check if we have the original HTML locally
            local_html_path = STORAGE_DIR / f"{deck_id}.html"
            if not local_html_path.exists():
                raise ValueError("Deck file not found")

        try:
            html_content = local_html_path.read_text(encoding="utf-8")
        except Exception as e:
            traceback.print_exc()
            raise RuntimeError(f"Failed to read deck HTML: {str(e)}")

        # Browser previews use authenticated application URLs for generated media.
        # Resolve only this deck's immutable media paths directly from S3 before
        # handing the SVG to resvg, which cannot authenticate against the app.
        media_path_pattern = re.compile(
            rf"/api/v1/ai/slides/{re.escape(deck_id)}/media/"
            r"([0-9a-f]{8}-[0-9a-f]{4}-[1-8][0-9a-f]{3}-"
            r"[89ab][0-9a-f]{3}-[0-9a-f]{12})",
            re.IGNORECASE,
        )
        for media_id in set(media_path_pattern.findall(html_content)):
            media_object = await download_bytes_from_s3(
                f"slides/{deck_id}/media/{media_id}.png"
            )
            if not media_object:
                raise RuntimeError(f"Failed to load generated slide media {media_id}")

            media_bytes, content_type = media_object
            if content_type not in {"image/png", "image/jpeg", "image/webp"}:
                raise RuntimeError(
                    f"Unsupported generated slide media type: {content_type}"
                )

            data_url = (
                f"data:{content_type};base64,"
                f"{base64.b64encode(media_bytes).decode('ascii')}"
            )
            media_url = f"/api/v1/ai/slides/{deck_id}/media/{media_id}"
            html_content = html_content.replace(media_url, data_url)

        # Extract all SVG markup
        svgs = re.findall(r"(<svg[^>]*>.*?</svg>)", html_content, re.DOTALL)
        if not svgs:
            raise ValueError("No SVG slides found in deck")

        # Generate PPTX in a specific output path
        pptx_path = STORAGE_DIR / f"{deck_id}.pptx"
        try:
            prs = Presentation()
            # Set 16:9 aspect ratio
            prs.slide_width = Emu(int(12192000))  # 13.33 in
            prs.slide_height = Emu(int(12192000 * 9 / 16))  # 7.5 in
            blank_layout = prs.slide_layouts[6]  # Blank slide layout

            def parse_svg_style(style_str: str | None) -> dict[str, str]:
                res = {}
                if not style_str:
                    return res
                for item in style_str.split(";"):
                    if ":" in item:
                        k, v = item.split(":", 1)
                        res[k.strip().lower()] = v.strip()
                return res

            def parse_svg_color(color_str: str | None) -> RGBColor | None:
                if not color_str or color_str.lower() in ("none", "transparent", "inherit"):
                    return None
                color_str = color_str.strip().lower()
                if color_str.startswith("#"):
                    c = color_str[1:]
                    if len(c) == 3:
                        c = "".join([x * 2 for x in c])
                    if len(c) == 6:
                        try:
                            return RGBColor(
                                int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
                            )
                        except ValueError:
                            pass
                m = re.match(r"rgba?\s*\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", color_str)
                if m:
                    return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                color_map = {
                    "white": RGBColor(255, 255, 255),
                    "black": RGBColor(0, 0, 0),
                    "red": RGBColor(255, 0, 0),
                    "blue": RGBColor(0, 0, 255),
                    "green": RGBColor(0, 128, 0),
                }
                return color_map.get(color_str)

            def parse_transform_translate(trans_str: str | None) -> tuple[float, float]:
                if not trans_str:
                    return 0.0, 0.0
                m = re.search(
                    r"translate\s*\(\s*([-0-9.]+)[,\s]+([-0-9.]+)\s*\)", trans_str
                )
                if m:
                    return float(m.group(1)), float(m.group(2))
                m1 = re.search(r"translate\s*\(\s*([-0-9.]+)\s*\)", trans_str)
                if m1:
                    return float(m1.group(1)), 0.0
                return 0.0, 0.0

            for idx, svg_markup in enumerate(svgs):
                clean_svg = re.sub(
                    r"&(?!(?:[a-zA-Z0-9]+|#[0-9]+|#x[0-9a-fA-F]+);)",
                    "&amp;",
                    svg_markup,
                )

                vb_match = re.search(r'viewBox=[\"\']([0-9.\s]+)[\"\']', clean_svg)
                if vb_match:
                    parts = [float(x) for x in vb_match.group(1).split()]
                    svg_w, svg_h = parts[2], parts[3]
                else:
                    svg_w, svg_h = 1440.0, 810.0

                try:
                    root = ET.fromstring(clean_svg)
                except ET.ParseError:
                    # Fallback to rendering whole slide as picture if XML parsing fails
                    png_bytes = bytes(
                        resvg_py.svg_to_bytes(svg_string=clean_svg, width=1920)
                    )
                    slide = prs.slides.add_slide(blank_layout)
                    slide.shapes.add_picture(
                        io.BytesIO(png_bytes), 0, 0, prs.slide_width, prs.slide_height
                    )
                    continue

                def remove_text_and_image_nodes(element: ET.Element) -> None:
                    to_remove = []
                    for child in list(element):
                        if child.tag.endswith("text") or child.tag.endswith("image"):
                            to_remove.append(child)
                        else:
                            remove_text_and_image_nodes(child)
                    for child in to_remove:
                        element.remove(child)

                bg_root = copy.deepcopy(root)
                remove_text_and_image_nodes(bg_root)
                bg_svg = ET.tostring(bg_root, encoding="utf-8").decode("utf-8")
                bg_png = bytes(
                    resvg_py.svg_to_bytes(svg_string=bg_svg, width=1920)
                )

                slide = prs.slides.add_slide(blank_layout)
                slide.shapes.add_picture(
                    io.BytesIO(bg_png), 0, 0, prs.slide_width, prs.slide_height
                )

                scale_x = prs.slide_width / svg_w
                scale_y = prs.slide_height / svg_h
                scale_pt = 540.0 / svg_h

                def extract_image_bytes(href: str | None) -> bytes | None:
                    if not href:
                        return None
                    if href.startswith("data:image/"):
                        try:
                            _, base64_str = href.split(",", 1)
                            return base64.b64decode(base64_str)
                        except Exception:
                            return None
                    return None

                def process_node(
                    node: ET.Element, parent_tx: float = 0.0, parent_ty: float = 0.0
                ) -> None:
                    tx, ty = parse_transform_translate(node.attrib.get("transform"))
                    curr_tx = parent_tx + tx
                    curr_ty = parent_ty + ty

                    if node.tag.endswith("text"):
                        style_dict = parse_svg_style(node.attrib.get("style"))

                        tspans = node.findall(".//{*}tspan")
                        lines = []
                        if tspans:
                            for ts in tspans:
                                t_text = (ts.text or "").strip()
                                if t_text:
                                    lines.append((t_text, ts))
                        else:
                            t_text = "".join(node.itertext()).strip()
                            if t_text:
                                lines.append((t_text, node))

                        if lines:
                            full_text = " ".join([l[0] for l in lines])
                            x = (
                                float(
                                    node.attrib.get("x", style_dict.get("x", 0))
                                )
                                + curr_tx
                            )
                            y = (
                                float(
                                    node.attrib.get("y", style_dict.get("y", 0))
                                )
                                + curr_ty
                            )

                            fs_str = (
                                node.attrib.get(
                                    "font-size", style_dict.get("font-size", "18")
                                )
                                .replace("px", "")
                                .replace("pt", "")
                            )
                            try:
                                font_size = float(fs_str)
                            except ValueError:
                                font_size = 18.0

                            font_weight = node.attrib.get(
                                "font-weight", style_dict.get("font-weight", "normal")
                            )
                            font_family = node.attrib.get(
                                "font-family", style_dict.get("font-family", "Arial")
                            )
                            fill_str = node.attrib.get(
                                "fill", style_dict.get("fill", "#000000")
                            )
                            text_anchor = node.attrib.get(
                                "text-anchor", style_dict.get("text-anchor", "start")
                            )
                            data_w_str = node.attrib.get(
                                "data-w", style_dict.get("data-w", "0")
                            )
                            data_w = float(data_w_str) if data_w_str else 0.0

                            box_w = (
                                data_w
                                if data_w > 0
                                else max(font_size * len(full_text) * 0.6, 200.0)
                            )
                            box_h = font_size * 1.3 * max(len(lines), 1)
                            top_svg = y - font_size * 0.85

                            if text_anchor == "middle":
                                left_svg = x - (box_w / 2.0)
                                align = PP_ALIGN.CENTER
                            elif text_anchor == "end":
                                left_svg = x - box_w
                                align = PP_ALIGN.RIGHT
                            else:
                                left_svg = x
                                align = PP_ALIGN.LEFT

                            left = Emu(int(left_svg * scale_x))
                            top = Emu(int(top_svg * scale_y))
                            width = Emu(int(box_w * scale_x))
                            height = Emu(int(box_h * scale_y))

                            txBox = slide.shapes.add_textbox(left, top, width, height)
                            tf = txBox.text_frame
                            tf.word_wrap = True
                            tf.margin_left = Inches(0.02)
                            tf.margin_right = Inches(0.02)
                            tf.margin_top = Inches(0.02)
                            tf.margin_bottom = Inches(0.02)

                            for line_idx, (line_text, line_elem) in enumerate(lines):
                                p = (
                                    tf.paragraphs[0]
                                    if line_idx == 0
                                    else tf.add_paragraph()
                                )
                                p.text = line_text
                                p.alignment = align
                                p.font.name = font_family.split(",")[0].strip(
                                    " \"'"
                                )
                                p.font.size = Pt(font_size * scale_pt)
                                p.font.bold = font_weight in (
                                    "bold",
                                    "700",
                                    "800",
                                    "900",
                                )

                                c = parse_svg_color(fill_str)
                                if c:
                                    p.font.color.rgb = c
                    elif node.tag.endswith("image"):
                        href = node.attrib.get("href") or node.attrib.get(
                            "{http://www.w3.org/1999/xlink}href"
                        )
                        if href:
                            img_bytes = extract_image_bytes(href)
                            if img_bytes:
                                try:
                                    style_dict = parse_svg_style(node.attrib.get("style"))
                                    x_str = str(node.attrib.get("x", style_dict.get("x", "0")))
                                    y_str = str(node.attrib.get("y", style_dict.get("y", "0")))
                                    w_str = str(node.attrib.get("width", style_dict.get("width", "0")))
                                    h_str = str(node.attrib.get("height", style_dict.get("height", "0")))

                                    x_val = float(x_str.replace("px", "")) + curr_tx
                                    y_val = float(y_str.replace("px", "")) + curr_ty
                                    w_val = float(w_str.replace("px", ""))
                                    h_val = float(h_str.replace("px", ""))

                                    if w_val > 0 and h_val > 0:
                                        img_left = Emu(int(x_val * scale_x))
                                        img_top = Emu(int(y_val * scale_y))
                                        img_width = Emu(int(w_val * scale_x))
                                        img_height = Emu(int(h_val * scale_y))

                                        slide.shapes.add_picture(
                                            io.BytesIO(img_bytes),
                                            img_left,
                                            img_top,
                                            img_width,
                                            img_height,
                                        )
                                except Exception as img_err:
                                    print(f"[PPTX] Warning parsing image shape: {img_err}")
                    else:
                        for child in list(node):
                            process_node(child, curr_tx, curr_ty)

                process_node(root)

            prs.save(str(pptx_path))
        except Exception as e:
            traceback.print_exc()
            raise RuntimeError(f"Failed to assemble PPTX: {str(e)}")

        return pptx_path


# --- template review: inspect extracted slots, and correct them ---------------
async def inspect_collection(collection: str) -> Dict[str, Any]:
    """Every category's detected slots + warnings, for a pre-save review screen.

    Extraction infers intent from geometry, so some decisions are wrong in ways
    only a human can spot. This exposes them: what each slot is called, what the
    planner is told to write there, how much room it has, and which slots look
    suspicious (huge box with a tiny budget, a 'title' big enough for prose,
    text that cannot wrap).
    """
    svc = SlideService()
    library_dir = await svc._ensure_collection_downloaded(collection)
    report = await run_in_threadpool(slide_skills.inspect_template, str(library_dir))
    return {
        "collection": collection,
        "categories": report,
        "warning_count": sum(len(c["warnings"]) for c in report),
    }


async def render_collection_overlay(collection: str, category: str,
                                    variant: str = "standard",
                                    boxes: bool = True,
                                    editable: bool = False) -> str:
    """The category's slide, optionally with every slot outlined and labelled.

    `boxes=False` returns the bare slide. `editable=True` additionally tags every
    element with the slot it belongs to and swaps {{placeholders}} for readable
    sample copy, so the browser can drag the REAL text instead of an empty
    outline over a static picture.
    """
    svc = SlideService()
    library_dir = await svc._ensure_collection_downloaded(collection)
    svg_path = Path(library_dir) / category / f"{variant}.svg"
    if not svg_path.exists():
        raise FileNotFoundError(f"{category}/{variant}.svg not found in {collection}")

    def _build() -> str:
        raw = svg_path.read_text(encoding="utf-8")
        if editable:
            info = slide_skills.inspect_variant(svg_path)
            return slide_skills.prepare_editable_svg(
                raw, info["slots"], PREVIEW_SAMPLE_DATA)
        if not boxes:
            return raw
        info = slide_skills.inspect_variant(svg_path)
        return slide_skills.render_slot_overlay(raw, info["slots"])

    return await run_in_threadpool(_build)


async def update_collection_slots(collection: str, category: str, variant: str,
                                  edits: list) -> Dict[str, Any]:
    """Apply reviewer corrections, then re-upload the changed files to S3."""
    svc = SlideService()
    library_dir = await svc._ensure_collection_downloaded(collection)
    result = await run_in_threadpool(
        slide_skills.apply_slot_edits, str(library_dir), category, variant, edits)

    # push the corrected SVG + schema back so later generations use them
    from app.services.s3_service import upload_file_to_s3

    bucket = svc._template_bucket(collection)
    synced_all = True
    for name in (f"{variant}.svg", f"{variant}.schema.json"):
        local = Path(library_dir) / category / name
        if local.exists():
            uploaded = await upload_file_to_s3(
                local, f"templates/{collection}/{category}/{name}",
                bucket_name=bucket)
            if not uploaded:
                synced_all = False
    if not synced_all:
        raise RuntimeError(f"Failed to sync updated template slots for {category} to S3")
    result["synced"] = True
    return result


async def delete_collection_category(collection: str, category: str) -> Dict[str, Any]:
    """Drop one layout from a collection, locally and in S3.

    Extraction turns every slide or master layout into a category, so a deck
    imported as a template brings along duplicates and dividers nobody wants to
    generate onto. Removing one here stops the planner ever selecting it.

    The last remaining layout is refused: a collection with no categories is not
    an empty collection, it is a broken one that fails at generation time.
    """
    # `category` reaches this from a URL path, so it must not be able to climb
    # out of the collection directory.
    if not category or category in (".", "..") or "/" in category or "\\" in category:
        raise ValueError(f"Invalid category name: {category!r}")

    svc = SlideService()
    library_dir = await svc._ensure_collection_downloaded(collection)
    target = Path(library_dir) / category
    if not target.exists() or not target.is_dir():
        raise FileNotFoundError(f"{category} not found in {collection}")

    siblings = [
        child
        for child in Path(library_dir).iterdir()
        if child.is_dir() and any(child.glob("*.svg"))
    ]
    if len(siblings) <= 1:
        raise ValueError(
            f"'{category}' is the only layout left in '{collection}' — "
            "delete the whole collection instead of emptying it"
        )

    from app.services.s3_service import delete_s3_prefix

    removed_keys = await delete_s3_prefix(
        f"templates/{collection}/{category}/",
        bucket_name=svc._template_bucket(collection),
    )
    if removed_keys < 0:
        raise RuntimeError(f"Failed to delete S3 objects for layout '{category}' from '{collection}'")

    await run_in_threadpool(shutil.rmtree, target)
    logger.info(
        f"Deleted layout '{category}' from '{collection}' "
        f"({removed_keys} S3 object(s) removed)"
    )
    invalidate_collections_cache()
    return {
        "collection": collection,
        "category": category,
        "deleted": True,
        "s3_objects_removed": removed_keys,
        "remaining": len(siblings) - 1,
    }
