import uuid
from pathlib import Path
from typing import Any, Dict
from fastapi import APIRouter, BackgroundTasks, File, UploadFile, HTTPException
from fastapi.responses import FileResponse

from app.schemas.slide_schema import GenReq, PlanGenReq
from app.services.slide_service import SlideService
from app.deps import STORAGE_DIR
from app.services.s3_service import upload_file_to_s3, download_file_from_s3

router = APIRouter(prefix="/slides", tags=["Slides"])
slide_service = SlideService()

# Global memory job database
jobs: Dict[str, Dict[str, Any]] = {}


async def execute_generation_job(job_id: str, req: GenReq, out_path: Path):
    try:
        jobs[job_id]["status"] = "running"
        result = await slide_service.generate_deck(
            topic=req.topic,
            collection=req.collection,
            output_path=out_path,
            palette=req.palette,
            language=req.language,
            animation=req.animation,
        )
        # Upload to S3
        s3_key = f"slides/{job_id}.html"
        uploaded = await upload_file_to_s3(out_path, s3_key)

        jobs[job_id]["status"] = "done"
        jobs[job_id]["result"] = {
            "deck_id": job_id,
            "slides": result.get("slides", []),
            "usage": result.get("usage", {}),
        }
        if uploaded:
            jobs[job_id]["result"]["s3_key"] = s3_key
    except Exception as e:
        jobs[job_id]["status"] = "error"
        jobs[job_id]["message"] = str(e)


async def execute_plan_generation_job(job_id: str, req: PlanGenReq, out_path: Path):
    try:
        jobs[job_id]["status"] = "running"
        plan_dict = {"title": req.title, "slides": [s.model_dump() for s in req.slides]}
        result = await slide_service.generate_deck_from_plan(
            plan=plan_dict,
            output_path=out_path,
            palette=req.palette,
            images=req.images,
            image_source=req.image_source,
        )
        # Upload to S3
        s3_key = f"slides/{job_id}.html"
        uploaded = await upload_file_to_s3(out_path, s3_key)

        jobs[job_id]["status"] = "done"
        jobs[job_id]["result"] = {
            "deck_id": job_id,
            "slides": result.get("slides", []),
            "usage": result.get("usage", {}),
            "warnings": result.get("warnings", []),
        }
        if uploaded:
            jobs[job_id]["result"]["s3_key"] = s3_key
    except Exception as e:
        jobs[job_id]["status"] = "error"
        jobs[job_id]["message"] = str(e)


@router.get("/templates/categories")
async def get_categories():
    try:
        return await slide_service.get_categories()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/templates/collections")
async def get_collections():
    try:
        return await slide_service.get_collections()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/generate")
async def generate(req: GenReq, background_tasks: BackgroundTasks):
    job_id = uuid.uuid4().hex[:12]
    out_path = STORAGE_DIR / f"{job_id}.html"
    jobs[job_id] = {
        "status": "queued",
        "result": None,
        "message": None,
    }
    background_tasks.add_task(execute_generation_job, job_id, req, out_path)
    return {"job_id": job_id, "status": "queued"}


@router.post("/generate-from-plan")
async def generate_from_plan(req: PlanGenReq, background_tasks: BackgroundTasks):
    job_id = uuid.uuid4().hex[:12]
    out_path = STORAGE_DIR / f"{job_id}.html"
    jobs[job_id] = {
        "status": "queued",
        "result": None,
        "message": None,
    }
    background_tasks.add_task(execute_plan_generation_job, job_id, req, out_path)
    return {"job_id": job_id, "status": "queued"}


@router.get("/jobs/{job_id}")
async def get_job_status(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    return jobs[job_id]


@router.get("/decks/{deck_id}")
async def get_deck(deck_id: str):
    file_path = STORAGE_DIR / f"{deck_id}.html"
    if not file_path.exists():
        s3_key = f"slides/{deck_id}.html"
        downloaded = await download_file_from_s3(s3_key, file_path)
        if not downloaded:
            raise HTTPException(status_code=404, detail="Deck file not found")
    return FileResponse(file_path, media_type="text/html")


@router.get("/decks/{deck_id}/pptx")
async def get_deck_pptx(deck_id: str):
    import io
    import re
    from pptx import Presentation
    from pptx.util import Emu
    import resvg_py

    # Always fetch latest HTML from S3 to make sure we have visual edits
    s3_key = f"slides/{deck_id}.html"
    local_html_path = STORAGE_DIR / f"{deck_id}_latest.html"

    # Try downloading from S3
    downloaded = await download_file_from_s3(s3_key, local_html_path)
    if not downloaded:
        # Fallback: check if we have the original HTML locally
        local_html_path = STORAGE_DIR / f"{deck_id}.html"
        if not local_html_path.exists():
            raise HTTPException(status_code=404, detail="Deck file not found")

    try:
        html_content = local_html_path.read_text(encoding="utf-8")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to read deck HTML: {str(e)}")

    # Extract all SVG markup
    svgs = re.findall(r"(<svg[^>]*>.*?</svg>)", html_content, re.DOTALL)
    if not svgs:
        raise HTTPException(status_code=400, detail="No SVG slides found in deck")

    # Generate PPTX in a specific output path
    pptx_path = STORAGE_DIR / f"{deck_id}.pptx"
    try:
        prs = Presentation()
        # Set 16:9 aspect ratio
        prs.slide_width = Emu(int(12192000)) # 13.33 in
        prs.slide_height = Emu(int(12192000 * 9 / 16)) # 7.5 in
        blank_layout = prs.slide_layouts[6] # Blank slide layout

        for idx, svg_markup in enumerate(svgs):
            png_bytes = bytes(resvg_py.svg_to_bytes(svg_string=svg_markup, width=1920))
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                io.BytesIO(png_bytes), 
                0, 0, 
                prs.slide_width, 
                prs.slide_height
            )

        prs.save(str(pptx_path))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to assemble PPTX: {str(e)}")

    return FileResponse(
        pptx_path, 
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=f"deck-{deck_id}.pptx"
    )



@router.post("/templates/import")
async def import_templates(file: UploadFile = File(...), name: str | None = None):
    filename = file.filename or ""
    if not (filename.lower().endswith(".zip") or filename.lower().endswith(".svg")):
        raise HTTPException(
            status_code=400,
            detail="Only ZIP archive files or SVG template files are supported",
        )
    try:
        file_bytes = await file.read()
        res = await slide_service.import_template_collection(file_bytes, filename, name)
        return {"status": "success", "imported": res}
    except ValueError as val_err:
        raise HTTPException(status_code=400, detail=str(val_err))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")
    finally:
        await file.close()
