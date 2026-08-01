import sys
from pathlib import Path

EXTERNAL_SERVICES_DIR = Path(__file__).resolve().parents[1]
if str(EXTERNAL_SERVICES_DIR) not in sys.path:
    sys.path.insert(0, str(EXTERNAL_SERVICES_DIR))

import unittest
from pptx import Presentation
from app.deps import STORAGE_DIR
from app.services.slide_service import SlideService


class EditablePptxTests(unittest.IsolatedAsyncioTestCase):
    async def test_generate_pptx_creates_editable_textboxes(self) -> None:
        service = SlideService()
        deck_id = "test_editable_deck"

        sample_html = """
        <!DOCTYPE html>
        <html>
        <body>
          <section class="slide">
            <svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1440 810">
              <rect width="1440" height="810" fill="#0A0E1A"/>
              <text x="720" y="132" font-size="24" font-weight="bold" fill="#00F0FF" text-anchor="middle">Intro Title</text>
              <text x="720" y="280" font-size="72" font-weight="bold" fill="#FFFFFF" text-anchor="middle">Editable Slide Header</text>
              <text x="720" y="470" font-size="20" fill="#CCCCCC" text-anchor="middle">Subtitle or description content</text>
            </svg>
          </section>
        </body>
        </html>
        """

        html_file = STORAGE_DIR / f"{deck_id}.html"
        html_file.write_text(sample_html.strip(), encoding="utf-8")

        pptx_path = Path("/tmp/dummy.pptx")
        try:
            pptx_path = await service.generate_pptx(deck_id)
            self.assertTrue(pptx_path.exists())

            prs = Presentation(str(pptx_path))
            self.assertEqual(len(prs.slides), 1)
            slide = prs.slides[0]

            # Slide should have background image shape + textboxes
            textboxes = [shape for shape in slide.shapes if shape.has_text_frame]
            self.assertEqual(len(textboxes), 3)

            extracted_text = [tb.text_frame.text for tb in textboxes]
            self.assertIn("Intro Title", extracted_text)
            self.assertIn("Editable Slide Header", extracted_text)
            self.assertIn("Subtitle or description content", extracted_text)
        finally:
            if html_file.exists():
                html_file.unlink()
            if pptx_path.exists():
                pptx_path.unlink()


if __name__ == "__main__":
    unittest.main()
